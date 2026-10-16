import io

import docx
import pytest
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.ingestion.chunking import MAX_TOKENS, OVERLAP_TOKENS, ChunkData, chunk_sections, count_tokens
from app.ingestion.parse import SCANNED_PDF_MESSAGE, ParseError, Section, looks_like_heading, parse
from app.ingestion.validate import DOCX, MD, PDF, FileRejected, detect_type


def make_docx() -> bytes:
    d = docx.Document()
    d.add_heading("Leave Policy", level=0)
    d.add_heading("Annual Leave", level=1)
    d.add_paragraph("Employees receive 25 days of annual leave per year.")
    d.add_heading("Probation", level=2)
    d.add_paragraph("Leave cannot be taken during the first 3 months of probation.")
    d.add_heading("Sick Leave", level=1)
    d.add_paragraph("Up to 10 paid sick days per year.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf(pages: list[list[str]]) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for lines in pages:
        y = 800
        for line in lines:
            c.drawString(50, y, line)
            y -= 18
        c.showPage()
    c.save()
    return buf.getvalue()


def test_docx_headings_become_sections_with_parent_path():
    chunks = chunk_sections(parse(make_docx(), DOCX))
    sections = [c.section for c in chunks]
    assert sections == ["Annual Leave", "Annual Leave › Probation", "Sick Leave"]
    assert "25 days" in chunks[0].text
    assert all(c.page is None for c in chunks)


def test_markdown_sections_and_title_skipped():
    md = b"# Code of Conduct\n\nIntro text.\n\n## Gifts\n\nNo gifts over $50.\n\n### Reporting\n\nTell HR.\n"
    chunks = chunk_sections(parse(md, MD))
    assert [c.section for c in chunks] == [None, "Gifts", "Gifts › Reporting"]


def test_pdf_pages_and_headings_carried_across_pages():
    pdf = make_pdf([
        ["1. Purpose", "This policy explains expected behaviour at work for all staff."],
        ["Continued text about behaviour that spans onto the second page here."],
    ])
    chunks = chunk_sections(parse(pdf, PDF))
    assert [(c.section, c.page) for c in chunks] == [("1. Purpose", 1), ("1. Purpose", 2)]


def test_scanned_pdf_rejected():
    blank = make_pdf([[], []])
    with pytest.raises(ParseError, match="No extractable text"):
        parse(blank, PDF)
    assert "scanned PDFs" in SCANNED_PDF_MESSAGE


def test_long_section_split_within_limits_with_overlap():
    sentences = [f"Rule number {i} says employees must follow procedure {i} carefully." for i in range(400)]
    section = Section("Big Section", None, [" ".join(sentences)])
    chunks = chunk_sections([section])
    assert len(chunks) > 1
    assert all(c.token_count <= MAX_TOKENS for c in chunks)
    assert [c.index for c in chunks] == list(range(len(chunks)))
    # The start of each chunk repeats the end of the previous one.
    for prev, nxt in zip(chunks, chunks[1:]):
        tail = prev.text[-200:]
        assert tail.split(". ")[-1][:40] in nxt.text[: OVERLAP_TOKENS * 8]


def test_paragraphs_are_not_merged_across_sections():
    chunks = chunk_sections([Section("A", None, ["alpha text"]), Section("B", None, ["beta text"])])
    assert [(c.section, c.text) for c in chunks] == [("A", "alpha text"), ("B", "beta text")]


def test_embed_text_includes_title_and_section():
    c = ChunkData(0, "except during probation", "Annual Leave › Probation", None, count_tokens("x"))
    assert c.embed_text("Leave Policy").startswith("Leave Policy — Annual Leave › Probation\n\n")
    assert ChunkData(0, "t", None, None, 1).embed_text("Doc").startswith("Doc\n\n")


@pytest.mark.parametrize("line,expected", [
    ("2.1 Sick Leave", True),
    ("ANNUAL LEAVE", True),
    ("Parental Leave and Family Care", True),
    ("Employees must notify their manager.", False),
    ("this is just a lowercase fragment", False),
])
def test_heading_heuristic(line, expected):
    assert looks_like_heading(line) is expected


def test_validation_by_content():
    assert detect_type(make_docx(), "x.docx") == DOCX
    assert detect_type(b"%PDF-1.7 ...", "x.pdf") == PDF
    with pytest.raises(FileRejected):
        detect_type(b"plain text, not a pdf", "fake.pdf")
    with pytest.raises(FileRejected):
        detect_type(b"\x00\x01binary", "notes.txt")
    with pytest.raises(FileRejected):
        detect_type(b"MZ...", "virus.exe")
