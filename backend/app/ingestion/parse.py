import io
import re
from dataclasses import dataclass, field

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from app.ingestion.validate import DOCX, MD, PDF


class ParseError(ValueError):
    """Raised with a message that is safe to show to the admin."""

SCANNED_PDF_MESSAGE = "No extractable text — scanned PDFs aren't supported yet."
# Fewer non-whitespace characters than this per page on average means the PDF
# is almost certainly images of text (a scan), not real text.
MIN_CHARS_PER_PAGE = 20


@dataclass
class Section:
    heading: str | None
    page: int | None = None
    paragraphs: list[str] = field(default_factory=list)


def parse(data: bytes, mime: str) -> list[Section]:
    if mime == PDF:
        sections = parse_pdf(data)
    elif mime == DOCX:
        sections = parse_docx(data)
    elif mime == MD:
        sections = parse_markdown(data.decode("utf-8-sig"))
    else:
        sections = parse_plaintext(data.decode("utf-8-sig"))
    return [s for s in sections if s.paragraphs]


_NUMBERED = re.compile(r"^\d+(\.\d+)*\.?\s+\S")
_SMALL_WORDS = {"a", "an", "and", "as", "at", "by", "for", "in", "of", "on", "or", "the", "to", "with", "&"}


def looks_like_heading(line: str) -> bool:
    line = line.strip()
    words = line.split()
    if not (2 < len(line) <= 80) or not words or len(words) > 10:
        return False
    if line[-1] in ".,;!?" or not any(c.isalpha() for c in line):
        return False
    if _NUMBERED.match(line):
        return True
    letters = [c for c in line if c.isalpha()]
    if len(letters) >= 3 and all(c.isupper() for c in letters):
        return True
    significant = [w for w in words if w.lower() not in _SMALL_WORDS]
    return bool(significant) and len(words) <= 8 and all(w[0].isupper() or not w[0].isalpha() for w in significant)


_LIST_ITEM = re.compile(r"^([-*•]|\d+[.)])\s+")


def _paragraphs_from_lines(lines: list[str]) -> list[str]:
    paras, current = [], ""
    for line in lines:
        line = line.strip()
        if not line:
            if current:
                paras.append(current)
            current = ""
        elif not current:
            current = line
        else:
            current += ("\n" if _LIST_ITEM.match(line) else " ") + line
    if current:
        paras.append(current)
    return paras


def parse_pdf(data: bytes) -> list[Section]:
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # pypdf raises a variety of errors on broken files
        raise ParseError(f"Could not read the PDF: {exc}") from exc

    total_chars = sum(len("".join(p.split())) for p in pages)
    if not pages or total_chars / len(pages) < MIN_CHARS_PER_PAGE:
        raise ParseError(SCANNED_PDF_MESSAGE)

    sections: list[Section] = []
    heading: str | None = None  # carried across page boundaries
    for page_no, text in enumerate(pages, start=1):
        # Every page starts a new section so each chunk maps to exactly one page.
        section = Section(heading, page_no)
        buffer: list[str] = []
        for line in text.splitlines():
            if looks_like_heading(line):
                section.paragraphs.extend(_paragraphs_from_lines(buffer))
                buffer = []
                sections.append(section)
                heading = line.strip()
                section = Section(heading, page_no)
            else:
                buffer.append(line)
        section.paragraphs.extend(_paragraphs_from_lines(buffer))
        sections.append(section)
    return sections


def _heading_level(p: Paragraph) -> int | None:
    name = (p.style.name if p.style is not None else "") or ""
    if name == "Title":
        return 0
    m = re.match(r"Heading (\d)", name)
    return int(m.group(1)) if m else None


def parse_docx(data: bytes) -> list[Section]:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"Could not read the DOCX file: {exc}") from exc

    # Heading path, e.g. ["Leave Types", "Sick Leave"] -> "Leave Types › Sick Leave",
    # so a sub-section keeps its parent's context.
    path: dict[int, str] = {}
    sections = [Section(None)]
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            level = _heading_level(block)
            if level is not None:
                if level == 0:  # document title: not a section heading
                    continue
                path = {lvl: h for lvl, h in path.items() if lvl < level}
                path[level] = text
                sections.append(Section(" › ".join(path[k] for k in sorted(path))))
            else:
                sections[-1].paragraphs.append(text)
        elif isinstance(block, Table):
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in block.rows]
            if rows:
                sections[-1].paragraphs.append("\n".join(rows))
    return sections


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")


def parse_markdown(text: str) -> list[Section]:
    path: dict[int, str] = {}
    sections = [Section(None)]
    buffer: list[str] = []
    for line in text.splitlines():
        m = _MD_HEADING.match(line)
        if m:
            sections[-1].paragraphs.extend(_paragraphs_from_lines(buffer))
            buffer = []
            level, heading = len(m.group(1)), m.group(2)
            if level == 1 and len(sections) == 1 and not sections[0].paragraphs:
                # A leading "# Title" is the document title, not a section.
                continue
            path = {lvl: h for lvl, h in path.items() if lvl < level}
            path[level] = heading
            sections.append(Section(" › ".join(path[k] for k in sorted(path))))
        else:
            buffer.append(line)
    sections[-1].paragraphs.extend(_paragraphs_from_lines(buffer))
    return sections


def parse_plaintext(text: str) -> list[Section]:
    sections = [Section(None)]
    buffer: list[str] = []
    for line in text.splitlines():
        if looks_like_heading(line):
            sections[-1].paragraphs.extend(_paragraphs_from_lines(buffer))
            buffer = []
            sections.append(Section(line.strip()))
        else:
            buffer.append(line)
    sections[-1].paragraphs.extend(_paragraphs_from_lines(buffer))
    return sections
