"""Generate the binary sample documents in backend/sample_docs/.

Run once (output is committed):  python scripts/make_sample_docs.py
Northwind Labs is a fictional company; all content is sample data.

The Employee Handbook (older) and Leave Policy 2026 (newer) deliberately
disagree on paid sick leave (10 vs 12 days), so the demo can show AskHR
flagging documents that contradict each other.
"""
from pathlib import Path

import docx
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).resolve().parent.parent / "sample_docs"

LEAVE_POLICY = [
    ("h1", "1. Purpose and Scope"),
    ("p", "This policy explains the types of leave available to Northwind Labs employees and how to request them. It applies to all full-time and part-time employees. Contractors and interns should refer to their contract."),
    ("h1", "2. Annual Leave"),
    ("h2", "2.1 Entitlement"),
    ("p", "All full-time employees are entitled to twenty (20) working days of paid annual leave in each leave year (1 January to 31 December). Leave accrues at 1.67 days per completed month of service. Part-time employees accrue leave pro rata to their contracted hours."),
    ("h2", "2.2 Probation"),
    ("p", "Leave accrues during the six-month probation period, but annual leave cannot be taken during probation except for up to two (2) days with written approval from your line manager."),
    ("h2", "2.3 Carry-Over"),
    ("p", "A maximum of five (5) unused days may be carried into the following leave year and must be taken by 31 March. Any balance remaining after this date is forfeited and will not be paid out."),
    ("h2", "2.4 Requesting Annual Leave"),
    ("p", "Request annual leave in the PeopleHub portal at least two weeks before the first day of leave. Your line manager approves or declines the request within five working days. Requests of more than ten consecutive working days also need approval from your head of department."),
    ("h1", "3. Public Holidays"),
    ("p", "Public holidays declared by the Federal Government of Nigeria are in addition to annual leave. Employees required to work on a public holiday receive a replacement day off within the following month."),
    ("h1", "4. Sick Leave"),
    ("h2", "4.1 Entitlement"),
    ("p", "Employees are entitled to twelve (12) working days of paid sick leave per leave year. A medical certificate from a registered doctor or your HMO hospital is required for absences longer than two consecutive working days."),
    ("h2", "4.2 Reporting Sickness"),
    ("p", "Tell your line manager before 9:00 on the first day of absence, by phone or WhatsApp, and record the absence in the PeopleHub portal when you return. Unused sick leave does not carry over."),
    ("h1", "5. Maternity Leave"),
    ("p", "Female employees with at least six months of continuous service receive sixteen (16) weeks of fully paid maternity leave. Notify People & Culture at least eight weeks before the expected start date. Nursing mothers may close one hour early each day for six months after returning to work."),
    ("h1", "6. Bereavement Leave"),
    ("p", "Employees may take up to five (5) days of paid bereavement leave on the death of a spouse, child, parent or sibling, and up to two (2) days for other relatives."),
    ("h1", "7. Questions"),
    ("p", "Questions about this policy should be directed to People & Culture at hr@northwindlabs.example."),
]

# (heading, paragraphs); "PAGE" starts a new page so citations land on varied pages.
HANDBOOK = [
    ("1. Welcome to Northwind", [
        "This handbook explains how we work at Northwind Labs: our hours, pay, benefits and the practical things every employee needs to know. Specific policies, such as the Leave Policy and the Code of Conduct, are published separately and take precedence where they go into more detail.",
    ]),
    ("2. Working Hours", [
        "Our standard working hours are 8:30 to 17:00, Monday to Friday, with a one-hour lunch break. Claims and customer support teams work a rota that includes some Saturday mornings; your manager will share the rota a month in advance.",
        "Northwind works in a hybrid model. Employees based in Lagos or Abuja are expected in the office on Tuesdays, Wednesdays and Thursdays. Monday and Friday are flexible days that you may work from home.",
    ]),
    ("3. Probation", [
        "New employees serve a six-month probation period. Your line manager will hold a check-in at three months and a confirmation review at six months.",
    ]),
    "PAGE",
    ("4. Pay and Payslips", [
        "Salaries are paid monthly on the 25th. If the 25th falls on a weekend or public holiday, salaries are paid on the previous working day. Payslips are available in the PeopleHub portal.",
        "Northwind deducts 8% of your monthly emoluments for your pension, and Northwind contributes a further 10%, paid to the Pension Fund Administrator you choose, in line with the Pension Reform Act.",
    ]),
    ("5. Health Insurance", [
        "Every employee is covered by Northwind Health, our own HMO plan, from their first day. Cover includes the employee, a spouse and up to four children under 18. It covers outpatient consultations, prescribed drugs, hospital admission, maternity care and basic dental and optical care. It does not cover cosmetic procedures.",
    ]),
    "PAGE",
    ("6. Time Off", [
        "Annual leave entitlements are set out in the Leave Policy.",
    ]),
    ("6.4 Requesting Time Off", [
        "Request time off in the PeopleHub portal and let your line manager know. Your manager approves requests in the portal. Please plan leave around month-end closing if you work in Finance or Claims.",
    ]),
    ("6.7 Sickness Absence", [
        "Staff may take up to ten (10) working days of paid sick leave in a calendar year. Further absence is unpaid unless agreed with People & Culture.",
    ]),
    "PAGE",
    ("7. Equipment and Access", [
        "Every employee receives a laptop on their first day. To get access to internal tools, such as the claims dashboard or the underwriting system, raise a request with the IT helpdesk at it@northwindlabs.example and copy your line manager, who approves it.",
    ]),
    ("8. Learning and Development", [
        "Each employee has a learning budget of 150,000 naira per year for courses, certifications and books related to their role. Employees studying for professional insurance exams (CIIN) receive three days of paid study leave per exam.",
    ]),
]

CODE_OF_CONDUCT = [
    ("h", "1. Purpose"),
    ("p", "This Code of Conduct sets out the standards of behaviour expected of everyone who works at Northwind Labs, including employees, contractors, agents and interns. It applies in the office, when working remotely, at company events and when representing Northwind externally."),
    ("h", "2. Respect and Anti-Harassment"),
    ("p", "We treat everyone with respect. Harassment, bullying, discrimination and intimidation of any kind are not tolerated, whether in person, in writing or online."),
    ("h", "3. Conflicts of Interest"),
    ("p", "A conflict of interest is any situation where your personal interests could influence, or appear to influence, decisions you make for Northwind. Examples include handling a claim for a friend or relative, holding shares in a supplier or competitor, outside employment with another insurer, and close personal relationships with people you manage. You must disclose any conflict to your line manager and People & Culture in writing, and step back from the related decision."),
    ("h", "4. Gifts and Hospitality"),
    ("p", "You may accept gifts or hospitality from customers, brokers or suppliers worth up to 50,000 naira. Anything above that value must be declared to your manager within five working days and may need to be returned. Cash, airtime or gift cards must never be accepted."),
    ("h", "5. Confidentiality and Customer Data"),
    ("p", "Policyholder data, claims records, pricing models and employee records are confidential. Handle personal data in line with the Nigeria Data Protection Act. Share it only with people who need it to do their job. Confidentiality obligations continue after you leave Northwind."),
    ("h", "6. Reporting Concerns"),
    ("p", "If you see something that breaks this Code, speak to your manager or People & Culture, or report it confidentially to ethics@northwindlabs.example. Reports can be made anonymously. Northwind does not tolerate retaliation against anyone who raises a concern in good faith."),
    ("h", "7. Consequences"),
    ("p", "Breaches of this Code may lead to disciplinary action, up to and including dismissal. Fraud, including claims fraud, will be reported to the relevant authorities."),
]

TRAVEL_POLICY = [
    ("h1", "1. Scope"),
    ("p", "This policy covers travel and out-of-pocket expenses incurred on Northwind business, including claims inspection visits, agent visits and travel between offices."),
    ("h1", "2. Claims Inspection Visits"),
    ("p", "Claims staff visiting a customer or a repair workshop within Lagos may claim actual ride-hailing or fuel costs. Keep the trip receipt or a screenshot of the ride."),
    ("h1", "3. Overnight Trips"),
    ("p", "For approved overnight trips outside your base city, Northwind pays for a standard hotel room booked through People & Culture, plus a per diem of 25,000 naira per night to cover meals and local transport. Flights must be economy class and booked at least seven days ahead where possible."),
    ("h1", "4. Mileage"),
    ("p", "If you use your own car for business travel, you may claim 150 naira per kilometre. Record the start and end odometer readings."),
    ("h1", "5. Submitting Expenses"),
    ("p", "Submit expenses in the PeopleHub portal within 30 days of the trip, with receipts attached. Your line manager approves them, and approved expenses are reimbursed with the next monthly salary."),
]


def make_leave_policy() -> None:
    d = docx.Document()
    d.add_heading("Leave Policy 2026", level=0)
    d.add_paragraph("Northwind Labs — Effective 1 January 2026 — Owner: People & Culture")
    for kind, text in LEAVE_POLICY:
        if kind == "h1":
            d.add_heading(text, level=1)
        elif kind == "h2":
            d.add_heading(text, level=2)
        else:
            d.add_paragraph(text)
    d.save(OUT / "leave_policy_2026.docx")


def make_travel_policy() -> None:
    d = docx.Document()
    d.add_heading("Travel & Expenses Policy", level=0)
    d.add_paragraph("Northwind Labs — Effective 1 March 2026 — Owner: Finance")
    for kind, text in TRAVEL_POLICY:
        if kind == "h1":
            d.add_heading(text, level=1)
        else:
            d.add_paragraph(text)
    d.save(OUT / "travel_expenses_policy.docx")


def make_handbook() -> None:
    styles = getSampleStyleSheet()
    story = [Paragraph("NORTHWIND LABS EMPLOYEE HANDBOOK", styles["Title"]), Spacer(1, 12)]
    for item in HANDBOOK:
        if item == "PAGE":
            story.append(PageBreak())
            continue
        heading, paras = item
        story.append(Paragraph(heading, styles["Heading2"]))
        for p in paras:
            story.append(Paragraph(p, styles["BodyText"]))
            story.append(Spacer(1, 8))
    SimpleDocTemplate(str(OUT / "employee_handbook.pdf"), pagesize=A4, title="Employee Handbook").build(story)


def make_code_of_conduct() -> None:
    styles = getSampleStyleSheet()
    story = [Paragraph("NORTHWIND LABS CODE OF CONDUCT", styles["Title"]), Spacer(1, 12)]
    for kind, text in CODE_OF_CONDUCT:
        story.append(Paragraph(text, styles["Heading2"] if kind == "h" else styles["BodyText"]))
        story.append(Spacer(1, 8))
    SimpleDocTemplate(str(OUT / "code_of_conduct.pdf"), pagesize=A4, title="Code of Conduct").build(story)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    make_leave_policy()
    make_travel_policy()
    make_handbook()
    make_code_of_conduct()
    print(f"Wrote sample documents to {OUT}")
